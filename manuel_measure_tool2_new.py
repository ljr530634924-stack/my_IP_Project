"""
manuel_measure_tool2_new.py

功能:
1. 手动在 ch00 上放置测量圆环，自动测量 ch01 圆环区域内的平均强度。
2. 逻辑：每4个圆环一组 (Q1->Q2->Q3->Q4)，循环颜色 (蓝->绿->紫->黄)。
3. 圆环由外圆半径与内外半径比例 (Percentage) 决定，测量值为环形区域内的平均强度。
4. 输出：生成包含各象限均值、面积及统计信息的 Excel 表格，以及带有 ID (位于四点中心) 的可视化图。
"""

import cv2
import numpy as np
import os
import sys
import glob
import tkinter as tk
from tkinter import filedialog, messagebox, Scrollbar, Canvas
from PIL import Image, ImageTk
import gc

# 尝试导入 openpyxl
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
except ImportError:
    print("Error: openpyxl library is missing. Please install it using 'pip install openpyxl'")

# --- Configuration ---
MEASURE_RADIUS = 70       # 测量外圆半径 (像素)
DEFAULT_PERCENTAGE = 0.5  # 内圆半径相对于外圆半径的默认比例

class ManualMeasurer2:
    def __init__(self):
        self.root = tk.Tk()
        self.root.title("Manual Measurement Tool 2 (4-Quadrant Annular Ring Mode)")
        self.root.geometry("1200x800")

        # Data
        self.points = []      # 存储选中的点 [(x, y), ...]
        self.raw_ch00 = None  # BGR numpy array (用于处理)
        self.vis_ch00 = None  # RGB numpy array (用于显示)

        self.tk_img = None    # 保持引用防止GC回收
        self.zoom_level = 1.0
        self.is_first_load = True
        self.ch00_path = ""
        self.ch01_path = ""
        self.image_files = []

        self.raw_w = 0
        self.raw_h = 0

        self.measure_radius = MEASURE_RADIUS      # 外圆半径
        self.percentage = DEFAULT_PERCENTAGE      # 内圆半径 / 外圆半径

        # 象限配置
        self.colors_tk = ["#0000FF", "#00FF00", "#FF00FF", "#FFFF00"]  # 蓝, 绿, 紫(洋红), 黄
        self.colors_cv = [(255, 0, 0), (0, 255, 0), (255, 0, 255), (0, 255, 255)]  # BGR
        self.quadrant_names = ["Q1 (Blue)", "Q2 (Green)", "Q3 (Purple)", "Q4 (Yellow)"]

        self._setup_ui()

    def _setup_ui(self):
        # 1. 工具栏
        toolbar = tk.Frame(self.root, bd=1, relief=tk.RAISED)
        toolbar.pack(side=tk.TOP, fill=tk.X)

        btn_load = tk.Button(toolbar, text="Load New Pair", command=self.load_images)
        btn_load.pack(side=tk.LEFT, padx=5, pady=5)

        tk.Label(toolbar, text="Zoom:").pack(side=tk.LEFT, padx=5)
        self.scale_zoom = tk.Scale(toolbar, from_=0.1, to=5.0, resolution=0.01, orient=tk.HORIZONTAL, command=self.on_zoom_change)
        self.scale_zoom.set(1.0)
        self.scale_zoom.pack(side=tk.LEFT, padx=5)

        tk.Label(toolbar, text="Outer Radius:").pack(side=tk.LEFT, padx=5)
        self.scale_radius = tk.Scale(toolbar, from_=10, to=200, orient=tk.HORIZONTAL, command=self.on_radius_change)
        self.scale_radius.set(self.measure_radius)
        self.scale_radius.pack(side=tk.LEFT, padx=5)

        tk.Label(toolbar, text="Inner/Outer (%):").pack(side=tk.LEFT, padx=5)
        self.scale_percentage = tk.Scale(toolbar, from_=0, to=99, orient=tk.HORIZONTAL, command=self.on_percentage_change)
        self.scale_percentage.set(int(self.percentage * 100))
        self.scale_percentage.pack(side=tk.LEFT, padx=5)

        btn_undo = tk.Button(toolbar, text="Undo (Right Click)", command=self.undo_point)
        btn_undo.pack(side=tk.LEFT, padx=5)

        btn_save = tk.Button(toolbar, text="Save & Export (Enter)", command=self.save_data, bg="#ddffdd")
        btn_save.pack(side=tk.LEFT, padx=20)

        # 状态栏
        self.lbl_status = tk.Label(toolbar, text="Please load images.", font=("Arial", 12, "bold"), width=40)
        self.lbl_status.pack(side=tk.RIGHT, padx=10)

        # 2. 主容器 (左右分栏)
        container = tk.Frame(self.root)
        container.pack(fill=tk.BOTH, expand=True)

        # 2.1 左侧文件列表
        frame_left = tk.Frame(container, width=200, bg="#f0f0f0")
        frame_left.pack(side=tk.LEFT, fill=tk.Y)

        sb_y = tk.Scrollbar(frame_left, orient=tk.VERTICAL)
        sb_x = tk.Scrollbar(frame_left, orient=tk.HORIZONTAL)

        self.lst_files = tk.Listbox(frame_left, width=30, font=("Arial", 10), selectmode=tk.SINGLE,
                                    yscrollcommand=sb_y.set, xscrollcommand=sb_x.set)

        sb_y.config(command=self.lst_files.yview)
        sb_x.config(command=self.lst_files.xview)

        sb_y.pack(side=tk.RIGHT, fill=tk.Y)
        sb_x.pack(side=tk.BOTTOM, fill=tk.X)
        self.lst_files.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.lst_files.bind("<<ListboxSelect>>", self.on_file_select)

        # 2.2 右侧画布区域
        frame_right = tk.Frame(container)
        frame_right.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True)

        self.v_scroll = Scrollbar(frame_right, orient=tk.VERTICAL)
        self.h_scroll = Scrollbar(frame_right, orient=tk.HORIZONTAL)

        self.canvas = Canvas(frame_right, bg="#202020",
                             xscrollcommand=self.h_scroll.set,
                             yscrollcommand=self.v_scroll.set)

        self.v_scroll.config(command=self.canvas.yview)
        self.h_scroll.config(command=self.canvas.xview)

        self.v_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        self.h_scroll.pack(side=tk.BOTTOM, fill=tk.X)
        self.canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        # 事件绑定
        self.canvas.bind("<Motion>", self.on_mouse_move)
        self.canvas.bind("<Button-1>", self.on_left_click)
        self.canvas.bind("<Button-3>", lambda e: self.undo_point())
        self.root.bind("<Return>", lambda e: self.save_data())
        self.root.bind("<space>", lambda e: self.save_data())

    def run(self):
        self.root.mainloop()

    def inner_radius(self):
        """计算内圆半径"""
        return max(0, int(self.measure_radius * self.percentage))

    def on_file_select(self, event):
        """列表选择事件：自动保存并加载新图"""
        selection = self.lst_files.curselection()
        if not selection:
            return

        index = selection[0]
        selected_file = self.image_files[index]

        if selected_file == self.ch00_path:
            return

        if self.points:
            print(f"Auto-saving before switching to {os.path.basename(selected_file)}...")
            self.save_data(silent=True)

        self.load_image_file(selected_file)

    def load_images(self):
        """加载图片按钮：打开文件并初始化列表"""
        file_path = filedialog.askopenfilename(
            title="Select ch00 image",
            filetypes=[("TIFF images", "*.tif"), ("All files", "*.*")]
        )
        if not file_path:
            return

        directory = os.path.dirname(file_path)

        search_pattern = os.path.join(directory, "*ch00*.tif")
        files = glob.glob(search_pattern)
        self.image_files = sorted(files)

        self.lst_files.delete(0, tk.END)
        for i, f in enumerate(self.image_files):
            self.lst_files.insert(tk.END, os.path.basename(f))
            base_name = os.path.splitext(os.path.basename(f))[0]
            result_path = os.path.join(directory, f"{base_name}_results.xlsx")
            if os.path.exists(result_path):
                self.lst_files.itemconfig(i, {'bg': '#ddffdd'})

        self.load_image_file(file_path)

        try:
            idx = self.image_files.index(file_path)
            self.lst_files.selection_clear(0, tk.END)
            self.lst_files.selection_set(idx)
            self.lst_files.see(idx)
        except ValueError:
            pass

    def load_from_directory(self, directory):
        """从指定目录自动加载文件列表并打开第一张图"""
        if not os.path.isdir(directory):
            return

        search_pattern = os.path.join(directory, "*ch00*.tif")
        files = glob.glob(search_pattern)
        self.image_files = sorted(files)

        if not self.image_files:
            messagebox.showinfo("Info", "No suitable images found in folder.")
            return

        self.lst_files.delete(0, tk.END)
        for i, f in enumerate(self.image_files):
            self.lst_files.insert(tk.END, os.path.basename(f))
            base_name = os.path.splitext(os.path.basename(f))[0]
            result_path = os.path.join(directory, f"{base_name}_results.xlsx")
            if os.path.exists(result_path):
                self.lst_files.itemconfig(i, {'bg': '#ddffdd'})

        self.load_image_file(self.image_files[0])

        self.lst_files.selection_clear(0, tk.END)
        self.lst_files.selection_set(0)
        self.lst_files.see(0)

    def load_image_file(self, file_path):
        """加载单张图片的实际逻辑"""
        self.raw_ch00 = None
        self.vis_ch00 = None
        self.tk_img = None
        gc.collect()

        self.ch00_path = file_path
        directory = os.path.dirname(file_path)
        filename = os.path.basename(file_path)

        if "ch00" in filename:
            ch01_path = os.path.join(directory, filename.replace("ch00", "ch01"))
        else:
            ch01_path = os.path.join(directory, filename.replace("ch00", "ch01"))

        if not os.path.exists(ch01_path):
            messagebox.showerror("Error", f"Could not find corresponding ch01 file:\n{ch01_path}")
            return

        self.ch01_path = ch01_path

        self.raw_ch00 = cv2.imread(self.ch00_path, cv2.IMREAD_UNCHANGED)
        if self.raw_ch00 is None:
            messagebox.showerror("Error", "Failed to load ch00 image.")
            return

        self.raw_h, self.raw_w = self.raw_ch00.shape[:2]
        self.points = []

        self.vis_ch00 = self.normalize_image(self.raw_ch00)

        if self.is_first_load:
            self.canvas.update_idletasks()
            canvas_w = self.canvas.winfo_width()
            canvas_h = self.canvas.winfo_height()

            if canvas_w > 10 and canvas_h > 10:
                scale_w = canvas_w / self.raw_w
                scale_h = canvas_h / self.raw_h
                initial_zoom = min(scale_w, scale_h) * 0.95
                self.scale_zoom.set(round(initial_zoom, 3))
            else:
                screen_h = self.root.winfo_screenheight()
                target_h = screen_h * 0.6
                initial_zoom = target_h / self.raw_h
                self.scale_zoom.set(round(initial_zoom, 2))

            self.is_first_load = False

        self.refresh_canvas()
        self.update_status_label()

    def normalize_image(self, img):
        """归一化任意位深图像到 8-bit RGB"""
        subsample = img[::10, ::10]
        vmin, vmax = np.percentile(subsample, (1, 99))

        if vmax > vmin:
            alpha = 255.0 / (vmax - vmin)
            beta = -vmin * alpha
        else:
            alpha = 0
            beta = 0

        img_8u = cv2.convertScaleAbs(img, alpha=alpha, beta=beta)

        if img_8u.ndim == 2:
            img_rgb = cv2.cvtColor(img_8u, cv2.COLOR_GRAY2RGB)
        else:
            img_rgb = cv2.cvtColor(img_8u, cv2.COLOR_BGR2RGB)
        return img_rgb

    def on_zoom_change(self, val):
        self.zoom_level = float(val)
        self.refresh_canvas()

    def on_radius_change(self, val):
        self.measure_radius = int(val)
        self.redraw_annotations()

    def on_percentage_change(self, val):
        self.percentage = int(val) / 100.0
        self.redraw_annotations()

    def refresh_canvas(self):
        if self.vis_ch00 is None:
            return

        new_w = int(self.raw_w * self.zoom_level)
        new_h = int(self.raw_h * self.zoom_level)

        resized_ch00 = cv2.resize(self.vis_ch00, (new_w, new_h), interpolation=cv2.INTER_NEAREST)

        pil_img = Image.fromarray(resized_ch00)
        self.tk_img = ImageTk.PhotoImage(pil_img)

        self.canvas.delete("all")
        self.canvas.create_image(0, 0, image=self.tk_img, anchor="nw")
        self.canvas.config(scrollregion=self.canvas.bbox("all"))

        self.redraw_annotations()

    def update_status_label(self):
        """更新状态栏，显示下一个圆环的颜色"""
        count = len(self.points)
        step = count % 4
        particle_idx = count // 4 + 1

        text = f"Next: {self.quadrant_names[step]} | Particle #{particle_idx} | Total Rings: {count}"

        bg_color = self.colors_tk[step]
        fg_color = "white" if step in [0, 2] else "black"

        self.lbl_status.config(text=text, bg=bg_color, fg=fg_color)

    def redraw_annotations(self):
        """重绘所有已放置的圆环"""
        self.canvas.delete("overlay")

        scale = self.zoom_level
        r_outer = self.measure_radius * scale
        r_inner = self.inner_radius() * scale

        for i, (px, py) in enumerate(self.points):
            cx = px * scale
            cy = py * scale

            step = i % 4
            color = self.colors_tk[step]

            # 外圆
            self.canvas.create_oval(cx - r_outer, cy - r_outer, cx + r_outer, cy + r_outer,
                                    outline=color, width=2, tags="overlay")
            # 内圆（虚线）
            if r_inner > 0:
                self.canvas.create_oval(cx - r_inner, cy - r_inner, cx + r_inner, cy + r_inner,
                                        outline=color, width=1, dash=(4, 4), tags="overlay")
            # 象限编号
            self.canvas.create_text(cx, cy, text=str(step + 1), fill=color,
                                    font=("Arial", 10, "bold"), tags="overlay")

    def on_mouse_move(self, event):
        if self.vis_ch00 is None:
            return

        canvas_x = self.canvas.canvasx(event.x)
        canvas_y = self.canvas.canvasy(event.y)

        self.canvas.delete("cursor")

        r_outer = self.measure_radius * self.zoom_level
        r_inner = self.inner_radius() * self.zoom_level

        step = len(self.points) % 4
        color = self.colors_tk[step]

        # 外圆预览
        self.canvas.create_oval(canvas_x - r_outer, canvas_y - r_outer,
                                canvas_x + r_outer, canvas_y + r_outer,
                                outline=color, width=1, dash=(4, 4), tags="cursor")
        # 内圆预览
        if r_inner > 0:
            self.canvas.create_oval(canvas_x - r_inner, canvas_y - r_inner,
                                    canvas_x + r_inner, canvas_y + r_inner,
                                    outline=color, width=1, dash=(4, 4), tags="cursor")

    def on_left_click(self, event):
        if self.vis_ch00 is None:
            return

        canvas_x = self.canvas.canvasx(event.x)
        canvas_y = self.canvas.canvasy(event.y)

        raw_x = int(canvas_x / self.zoom_level)
        raw_y = int(canvas_y / self.zoom_level)

        raw_x = max(0, min(raw_x, self.raw_w - 1))
        raw_y = max(0, min(raw_y, self.raw_h - 1))

        self.points.append((raw_x, raw_y))
        self.redraw_annotations()
        self.update_status_label()

    def undo_point(self):
        if self.points:
            self.points.pop()
            self.redraw_annotations()
            self.update_status_label()

    def save_data(self, silent=False):
        if not self.points:
            if not silent:
                messagebox.showwarning("Warning", "No rings placed.")
            return False

        # 检查分组完整性
        num_points = len(self.points)
        remainder = num_points % 4
        valid_count = num_points - remainder

        if valid_count == 0:
            if not silent:
                messagebox.showwarning("Warning", "Not enough rings to form a complete particle (need 4).")
            return False

        if remainder > 0:
            messagebox.showinfo("Incomplete Particle", f"Ignored the last {remainder} incomplete rings.")

        final_points = self.points[:valid_count]

        print("Calculating and saving...")

        if not os.path.exists(self.ch01_path):
            if not silent:
                messagebox.showerror("Error", "ch01 file missing.")
            return False
        raw_ch01 = cv2.imread(self.ch01_path, cv2.IMREAD_UNCHANGED)
        if raw_ch01 is None:
            if not silent:
                messagebox.showerror("Error", "Failed to load ch01.")
            return False

        r_outer = self.measure_radius
        r_inner = self.inner_radius()
        pct = int(self.percentage * 100)

        particle_results = []

        vis_mc = cv2.cvtColor(self.vis_ch00, cv2.COLOR_RGB2BGR)

        num_particles = valid_count // 4

        for i in range(num_particles):
            p_id = i + 1
            group = final_points[i * 4: (i + 1) * 4]

            means = []
            areas = []

            sum_x, sum_y = 0, 0

            for step, (cx, cy) in enumerate(group):
                sum_x += cx
                sum_y += cy

                # 构建环形 mask
                mask_outer = np.zeros(raw_ch01.shape[:2], dtype=np.uint8)
                cv2.circle(mask_outer, (cx, cy), r_outer, 255, -1)

                if r_inner > 0:
                    mask_inner = np.zeros(raw_ch01.shape[:2], dtype=np.uint8)
                    cv2.circle(mask_inner, (cx, cy), r_inner, 255, -1)
                    annulus_mask = cv2.subtract(mask_outer, mask_inner)
                else:
                    annulus_mask = mask_outer

                mean_val = cv2.mean(raw_ch01, mask=annulus_mask)[0]
                area_val = cv2.countNonZero(annulus_mask)

                means.append(mean_val)
                areas.append(area_val)

                # 可视化：外圆实线 + 内圆细线
                color = self.colors_cv[step]
                cv2.circle(vis_mc, (cx, cy), r_outer, color, 4)
                if r_inner > 0:
                    cv2.circle(vis_mc, (cx, cy), r_inner, color, 2)

            # 在几何中心画 ID
            center_x = int(sum_x / 4)
            center_y = int(sum_y / 4)
            cv2.putText(vis_mc, str(p_id), (center_x, center_y),
                        cv2.FONT_HERSHEY_SIMPLEX, 2.0, (255, 255, 255), 4)

            # Row: [id, Q1_m, Q2_m, Q3_m, Q4_m, Q1_a, Q2_a, Q3_a, Q4_a]
            row_data = [p_id] + means + areas
            particle_results.append(row_data)

        del raw_ch01
        gc.collect()

        # 统计 Footer
        stats_means = ["Mean"]
        stats_sd = ["SD"]
        stats_se = ["Std. Error"]

        num_cols = 8  # 4 Means + 4 Areas
        for col_idx in range(1, num_cols + 1):
            vals = [r[col_idx] for r in particle_results]
            if vals:
                m = np.mean(vals)
                s = np.std(vals, ddof=1) if len(vals) > 1 else 0
                se = s / np.sqrt(len(vals)) if len(vals) > 0 else 0
                stats_means.append(m)
                stats_sd.append(s)
                stats_se.append(se)
            else:
                stats_means.append(0)
                stats_sd.append(0)
                stats_se.append(0)

        # 保存文件
        directory = os.path.dirname(self.ch00_path)
        base_name = os.path.splitext(os.path.basename(self.ch00_path))[0]

        # Excel
        excel_path = os.path.join(directory, f"{base_name}_results.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Manual Data"

        headers = [
            "particle_id",
            f"Q1_Annulus_Mean (outer={r_outer}, inner={r_inner}, {pct}%)",
            f"Q2_Annulus_Mean (outer={r_outer}, inner={r_inner}, {pct}%)",
            f"Q3_Annulus_Mean (outer={r_outer}, inner={r_inner}, {pct}%)",
            f"Q4_Annulus_Mean (outer={r_outer}, inner={r_inner}, {pct}%)",
            "Q1_Annulus_Area",
            "Q2_Annulus_Area",
            "Q3_Annulus_Area",
            "Q4_Annulus_Area"
        ]
        ws.append(headers)

        col_map = {
            2: "0000FF", 6: "0000FF",  # Q1
            3: "008000", 7: "008000",  # Q2
            4: "800080", 8: "800080",  # Q3
            5: "CCCC00", 9: "CCCC00"   # Q4
        }

        for cell in ws[1]:
            c_code = col_map.get(cell.column)
            if c_code:
                cell.font = Font(bold=True, color=c_code)
            else:
                cell.font = Font(bold=True)

        for row in particle_results:
            ws.append(row)

        ws.append(stats_means)
        ws.append(stats_sd)
        ws.append(stats_se)

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                c_code = col_map.get(cell.column)
                if c_code:
                    cell.font = Font(color=c_code)

        wb.save(excel_path)

        # Image
        vis_path = os.path.join(directory, f"{base_name}_00visualization.png")
        cv2.imwrite(vis_path, vis_mc)

        print(f"Saved {num_particles} particles.")
        print(f"Excel: {excel_path}")
        print(f"Vis: {vis_path}")

        try:
            if self.ch00_path in self.image_files:
                idx = self.image_files.index(self.ch00_path)
                self.lst_files.itemconfig(idx, {'bg': '#ddffdd'})
        except ValueError:
            pass

        if not silent:
            messagebox.showinfo("Success", f"Saved {num_particles} particles successfully!\nIgnored {remainder} rings.")
        else:
            self.lbl_status.config(text=f"Auto-saved {num_particles} particles.", bg="#ddffdd", fg="black")

        return True

if __name__ == "__main__":
    app = ManualMeasurer2()
    if len(sys.argv) > 1:
        app.load_from_directory(sys.argv[1])
    app.run()
