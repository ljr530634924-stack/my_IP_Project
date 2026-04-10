import cv2
import numpy as np
import os
import glob
import csv
from scipy import ndimage
from scipy.spatial import distance as sp_dist
from skimage import morphology

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False
    print("[WARN] openpyxl not found. Output will be CSV instead of XLSX.")

# --- Configuration ---
INPUT_FOLDER = r"D:\Ingenieurpraixs\test_18032026"

ROI_DIAMETER_RATIO = 0.85

SAVE_DEBUG_IMAGES = True

# [NEW] Background normalization: sigma as a fraction of the image's shorter side.
# Larger value = smoother background estimate (captures more low-frequency variation).
# Recommended range: 0.08 - 0.20. Start with 0.12 and adjust if needed.
BG_SIGMA_RATIO = 0.12

# Signal detection parameters (applied after normalization + stretch)
THRESHOLD_RATIO = 0.55    # [Lowered from 0.6] After normalization background is more uniform,
                           # a slightly lower threshold reveals faint signals more reliably.
MIN_AREA = 1000
MAX_AREA = 19000
MIN_CIRCULARITY = 0.4
SMOOTHING_RADIUS = 3

# Measurement parameters (applied on original raw image)
MEASURE_RADIUS = 70

# Post-processing filters
REMOVE_OVERLAPPING = False
ISOLATION_THRESHOLD_RATIO = 3

# Visualization
FONT_SCALE = 0.8


def get_roi_mask(shape, ratio):
    h, w = shape[:2]
    mask = np.zeros((h, w), dtype=np.uint8)
    center = (w // 2, h // 2)
    if ratio is None:
        mask[:] = 255
        radius = 0
    else:
        radius = int(min(h, w) * ratio / 2)
        cv2.circle(mask, center, radius, 255, -1)
    return mask, center, radius


def calculate_circularity(area, perimeter):
    if perimeter == 0:
        return 0
    return 4 * np.pi * area / (perimeter * perimeter)


def normalize_background(img_float, sigma):
    """
    Remove non-uniform illumination by dividing by a large-scale background estimate.

    The background is estimated using a large Gaussian blur, which captures only the
    low-frequency illumination envelope (e.g., bright petri dish edges, vignetting).
    Dividing by this estimate yields an image where local contrast is uniform
    regardless of absolute brightness position.

    Returns a float32 image. Values are scaled so that the mean background level
    is preserved (avoids very large or very small numbers downstream).
    """
    background = cv2.GaussianBlur(img_float, (0, 0), sigmaX=sigma, sigmaY=sigma)
    # Prevent division by zero or near-zero background
    background = np.maximum(background, 1.0)
    mean_bg = np.mean(background)
    normalized = (img_float / background) * mean_bg
    return normalized


def stretch_to_uint8(img_float, roi_mask, stretch_low=2.0, stretch_high=99.5):
    """
    Stretch image to 0-255 uint8 using percentiles computed ONLY on ROI pixels.

    This is the key fix for biased stretching: bright regions outside the ROI
    (e.g., dish edges) are excluded from the percentile calculation, so the
    stretch is driven entirely by the interior signal of interest.
    """
    roi_pixels = img_float[roi_mask > 0]
    if len(roi_pixels) == 0:
        return np.clip(img_float, 0, 255).astype(np.uint8)

    p_lo = np.percentile(roi_pixels, stretch_low)
    p_hi = np.percentile(roi_pixels, stretch_high)

    if p_hi > p_lo:
        stretched = (img_float - p_lo) * (255.0 / (p_hi - p_lo))
    else:
        stretched = img_float.copy()

    return np.clip(stretched, 0, 255).astype(np.uint8)


def process_image(ch01_path):
    directory = os.path.dirname(ch01_path)
    filename = os.path.basename(ch01_path)
    name_no_ext = os.path.splitext(filename)[0]

    print(f"Processing: {filename}")

    # --- Step 1: Read raw image ---
    raw_img = cv2.imread(ch01_path, cv2.IMREAD_UNCHANGED)
    if raw_img is None:
        print(f"  [ERROR] Cannot read: {ch01_path}")
        return None, None, None

    img_float = raw_img.astype(np.float32)

    # --- Step 2: ROI mask ---
    roi_mask, roi_center, roi_radius = get_roi_mask(img_float.shape, ROI_DIAMETER_RATIO)

    # --- Step 3: Background normalization (on full image) ---
    # We normalize the full image so that the Gaussian blur is not affected by
    # artificially zeroed-out border pixels. The ROI is applied later.
    h, w = img_float.shape[:2]
    sigma = max(int(min(h, w) * BG_SIGMA_RATIO), 15)
    print(f"  [INFO] Background normalization sigma = {sigma}px "
          f"({BG_SIGMA_RATIO*100:.0f}% of {min(h,w)}px short side)")

    normalized = normalize_background(img_float, sigma)

    # --- Step 4: Stretch to 8-bit using only ROI pixels for percentile calculation ---
    adj_img = stretch_to_uint8(normalized, roi_mask, stretch_low=2.0, stretch_high=99.5)

    if SAVE_DEBUG_IMAGES:
        cv2.imwrite(os.path.join(directory, f"{name_no_ext}_debug_01_enhanced.png"), adj_img)

    # --- Step 5: Threshold ---
    threshold_val = int(THRESHOLD_RATIO * 255)
    _, binary = cv2.threshold(adj_img, threshold_val, 255, cv2.THRESH_BINARY)

    # --- Step 6: Morphological cleanup ---
    binary_bool = morphology.remove_small_objects(binary > 0, min_size=MIN_AREA)
    binary_cleaned = binary_bool.astype(np.uint8) * 255
    binary_filled = (ndimage.binary_fill_holes(binary_cleaned > 0) * 255).astype(np.uint8)

    if SMOOTHING_RADIUS > 0:
        k_size = SMOOTHING_RADIUS * 2 + 1
        kernel = cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (k_size, k_size))
        binary_filled = cv2.morphologyEx(binary_filled, cv2.MORPH_OPEN, kernel)

    if SAVE_DEBUG_IMAGES:
        cv2.imwrite(os.path.join(directory, f"{name_no_ext}_debug_03_signal_mask_filled.png"), binary_filled)

    # --- Step 7: Apply ROI ---
    binary_roi = cv2.bitwise_and(binary_filled, binary_filled, mask=roi_mask)

    # --- Step 8: Find contours ---
    contours, _ = cv2.findContours(binary_roi, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    print(f"  [DEBUG] Found {len(contours)} contours before filtering.")

    # --- Step 9: Visualization setup ---
    vis_img_01 = cv2.cvtColor(adj_img, cv2.COLOR_GRAY2BGR)
    if ROI_DIAMETER_RATIO is not None:
        cv2.circle(vis_img_01, roi_center, roi_radius, (255, 255, 0), 2)

    ch00_path = ch01_path.replace("ch01", "ch00")
    vis_img_00 = None
    if os.path.exists(ch00_path):
        ch00_raw = cv2.imread(ch00_path, cv2.IMREAD_UNCHANGED)
        if ch00_raw is not None:
            if ch00_raw.dtype == np.uint16:
                vis_img_00 = (ch00_raw / 256).astype(np.uint8)
            else:
                vis_img_00 = ch00_raw.astype(np.uint8)
            vis_img_00 = cv2.cvtColor(vis_img_00, cv2.COLOR_GRAY2BGR)
            if ROI_DIAMETER_RATIO is not None:
                cv2.circle(vis_img_00, roi_center, roi_radius, (255, 255, 0), 2)

    # --- Step 10: Filter contours and collect measurements ---
    measurements = []

    if SAVE_DEBUG_IMAGES:
        debug_filtered = np.zeros_like(binary_roi)
        cv2.drawContours(debug_filtered, contours, -1, 255, -1)
        cv2.imwrite(os.path.join(directory, f"{name_no_ext}_debug_04_signal_mask_filtered.png"), debug_filtered)

    for cnt in contours:
        area = cv2.contourArea(cnt)
        if area < MIN_AREA:
            continue
        if area > MAX_AREA:
            cv2.drawContours(vis_img_01, [cnt], -1, (0, 0, 255), 1)
            continue

        perimeter = cv2.arcLength(cnt, True)
        circularity = calculate_circularity(area, perimeter)

        if circularity < MIN_CIRCULARITY:
            cv2.drawContours(vis_img_01, [cnt], -1, (0, 165, 255), 1)
            continue

        M = cv2.moments(cnt)
        if M["m00"] == 0:
            continue
        cx = int(M["m10"] / M["m00"])
        cy = int(M["m01"] / M["m00"])

        # Measure mean intensity on original raw image
        mask_c = np.zeros(raw_img.shape[:2], dtype=np.uint8)
        cv2.circle(mask_c, (cx, cy), MEASURE_RADIUS, 255, -1)
        mean_val = cv2.mean(raw_img, mask=mask_c)[0]

        measurements.append({
            "cx": cx, "cy": cy,
            "area": area,
            "circularity": circularity,
            "mean_intensity": mean_val
        })

    print(f"  -> {len(measurements)} candidate signals before post-processing.")
    final_measurements = measurements

    # --- Step 11: Post-processing filters ---
    if REMOVE_OVERLAPPING and len(final_measurements) > 1:
        points = np.array([[m['cx'], m['cy']] for m in final_measurements])
        dist_matrix = sp_dist.squareform(sp_dist.pdist(points))
        np.fill_diagonal(dist_matrix, np.inf)
        overlapping_pairs = np.argwhere(dist_matrix < 2 * MEASURE_RADIUS)
        indices_to_remove = np.unique(overlapping_pairs.flatten())
        if len(indices_to_remove) > 0:
            keep = np.ones(len(final_measurements), dtype=bool)
            keep[indices_to_remove] = False
            final_measurements = [m for i, m in enumerate(final_measurements) if keep[i]]
            print(f"  -> Removed {len(indices_to_remove)} overlapping signals.")

    if ISOLATION_THRESHOLD_RATIO > 0 and len(final_measurements) > 1:
        points = np.array([[m['cx'], m['cy']] for m in final_measurements])
        dist_matrix = sp_dist.squareform(sp_dist.pdist(points))
        np.fill_diagonal(dist_matrix, np.inf)
        min_distances = np.min(dist_matrix, axis=1)
        isolation_threshold = ISOLATION_THRESHOLD_RATIO * MEASURE_RADIUS
        isolated = min_distances > isolation_threshold
        num_isolated = np.sum(isolated)
        if num_isolated > 0:
            final_measurements = [m for i, m in enumerate(final_measurements) if not isolated[i]]
            print(f"  -> Removed {num_isolated} isolated signals.")

    # --- Step 12: Assign IDs and draw final visualization ---
    for i, m in enumerate(final_measurements):
        m['id'] = i + 1
        cx, cy = m['cx'], m['cy']
        cv2.circle(vis_img_01, (cx, cy), MEASURE_RADIUS, (0, 255, 0), 2)
        cv2.putText(vis_img_01, str(m['id']), (cx, cy),
                    cv2.FONT_HERSHEY_SIMPLEX, FONT_SCALE, (0, 255, 0), 2)
        if vis_img_00 is not None:
            cv2.circle(vis_img_00, (cx, cy), MEASURE_RADIUS, (0, 255, 0), 2)
            cv2.putText(vis_img_00, str(m['id']), (cx, cy),
                        cv2.FONT_HERSHEY_SIMPLEX, FONT_SCALE, (0, 255, 0), 2)

    vis_path_01 = None
    if SAVE_DEBUG_IMAGES:
        vis_path_01 = os.path.join(directory, f"{name_no_ext}_new_vis_01.png")
        cv2.imwrite(vis_path_01, vis_img_01)

    vis_path_00 = None
    if SAVE_DEBUG_IMAGES and vis_img_00 is not None:
        vis_name_00 = name_no_ext.replace("_ch01", "")
        vis_path_00 = os.path.join(directory, f"{vis_name_00}_new_vis_00.png")
        cv2.imwrite(vis_path_00, vis_img_00)

    return final_measurements, vis_path_01, vis_path_00


def save_results(measurements, output_path):
    if not measurements:
        return

    all_means = [m["mean_intensity"] for m in measurements]
    global_avg = np.mean(all_means) if all_means else 0
    global_std = np.std(all_means) if all_means else 0

    headers = ["Signal_ID", "Centroid_X", "Centroid_Y", "Contour_Area", "Circularity", "Mean_Intensity"]

    if HAS_OPENPYXL:
        wb = Workbook()
        ws = wb.active
        ws.title = "Signal Data"
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for m in measurements:
            ws.append([m["id"], m["cx"], m["cy"], m["area"], m["circularity"], m["mean_intensity"]])
        ws.append([])
        ws.append(["Summary Statistics"])
        ws.append(["Total Signals", len(measurements)])
        ws.append(["Global Average Intensity", global_avg])
        ws.append(["Std Dev", global_std])
        for i in range(4):
            ws[f"A{ws.max_row - i}"].font = Font(bold=True)
        wb.save(output_path)
    else:
        with open(output_path, 'w', newline='') as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            for m in measurements:
                writer.writerow([m["id"], m["cx"], m["cy"],
                                  m["area"], m["circularity"], m["mean_intensity"]])
            writer.writerow([])
            writer.writerow(["Summary Statistics"])
            writer.writerow(["Global Average Intensity", global_avg])


def main():
    print(f"=== DSA with Background Normalization | Folder: '{INPUT_FOLDER}' ===")
    print(f"    BG_SIGMA_RATIO={BG_SIGMA_RATIO}  THRESHOLD_RATIO={THRESHOLD_RATIO}  ROI={ROI_DIAMETER_RATIO}")

    search_pattern = os.path.join(INPUT_FOLDER, "*ch01*.tif")
    files = glob.glob(search_pattern)

    if not files:
        print("No ch01 files found.")
        return

    print(f"Found {len(files)} file(s).\n")

    for f in files:
        results, vis_path_01, vis_path_00 = process_image(f)

        if results:
            base_name = os.path.splitext(os.path.basename(f))[0]
            out_ext = "xlsx" if HAS_OPENPYXL else "csv"
            out_path = os.path.join(os.path.dirname(f), f"{base_name}_new_results.{out_ext}")
            save_results(results, out_path)
            avg = np.mean([m['mean_intensity'] for m in results])
            print(f"  -> Count={len(results)}, Avg Intensity={avg:.2f}")
            print(f"  -> Results: {os.path.basename(out_path)}")
            if vis_path_01:
                print(f"  -> Vis ch01:  {os.path.basename(vis_path_01)}")
            if vis_path_00:
                print(f"  -> Vis ch00:  {os.path.basename(vis_path_00)}")
        else:
            print("  -> No valid signals found.")
        print()


if __name__ == "__main__":
    main()
